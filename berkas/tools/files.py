from __future__ import annotations

import csv
import hashlib
import json
import os
import re

from berkas.entities import parse_text
from berkas.tools.base import ToolResult, edge, entity, evidence, fail, ok

try:
    from PIL import Image
    from PIL.ExifTags import GPSTAGS, TAGS
except ImportError:
    Image = None
    TAGS = {}
    GPSTAGS = {}

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


def sha256_file(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def detect_file_type(path: str, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    with open(path, "rb") as handle:
        head = handle.read(16)
    if (
        head.startswith(b"\xff\xd8\xff")
        or head.startswith(b"\x89PNG")
        or head.startswith(b"GIF87a")
        or head.startswith(b"GIF89a")
        or head.startswith(b"BM")
        or head.startswith(b"RIFF")
    ):
        return "image"
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"PK"):
        if ext == ".docx":
            return "docx"
        if ext in {".xlsx", ".xlsm"}:
            return "xlsx"
        return "archive"
    if ext in {".csv", ".tsv"}:
        return "csv"
    if ext == ".json":
        return "json"
    if ext in {".txt", ".md", ".log"}:
        return "text"
    return "unknown"


def _dms_to_decimal(values, ref: str) -> float:
    degrees, minutes, seconds = values
    decimal = float(degrees) + float(minutes) / 60 + float(seconds) / 3600
    return -decimal if ref in {"S", "W"} else decimal


def _entities_from_text(text: str, source: str) -> list:
    return [
        entity(
            item.type,
            item.value,
            item.normalized,
            item.confidence,
            source=source,
            **item.attributes,
        )
        for item in parse_text(text)
    ]


def analyze_image(path: str, filename: str) -> ToolResult:
    if Image is None:
        return fail("dependency_missing", "Pillow tidak terpasang.")
    try:
        image = Image.open(path)
    except Exception as exc:
        return fail("invalid_file", f"Gambar tidak bisa dibuka: {exc}")
    digest = sha256_file(path)
    file_entity = entity("image", filename, digest, 0.9, source="upload", filename=filename, size=os.path.getsize(path))
    entities = [file_entity]
    edges = []
    ev = [evidence("file", 1.0, "", f"File gambar {filename}, SHA256 {digest[:16]}", entity_key=file_entity.key)]
    warnings = []
    data = {
        "filename": filename,
        "sha256": digest,
        "format": image.format,
        "mode": image.mode,
        "width": image.width,
        "height": image.height,
        "exif": {},
        "gps": None,
        "ocr": None,
        "reverse_search": {
            "google_lens": "https://lens.google.com/uploadbyurl?url=",
            "tineye": "https://tineye.com/search",
            "yandex": "https://yandex.com/images/search",
        },
    }
    exif_raw = {}
    try:
        exif_raw = image.getexif() or {}
    except Exception:
        warnings.append("EXIF tidak bisa dibaca.")
    gps_data = {}
    for tag_id, value in exif_raw.items():
        tag = TAGS.get(tag_id, str(tag_id))
        if tag == "GPSInfo":
            for gps_id, gps_value in value.items():
                gps_data[GPSTAGS.get(gps_id, str(gps_id))] = gps_value
            continue
        if tag in {"Make", "Model", "Software", "Artist", "Copyright", "DateTimeOriginal", "BodySerialNumber", "LensModel"}:
            data["exif"][tag] = str(value)[:200]
    if "GPSLatitude" in gps_data and "GPSLongitude" in gps_data:
        try:
            lat = _dms_to_decimal(gps_data["GPSLatitude"], gps_data.get("GPSLatitudeRef", "N"))
            lon = _dms_to_decimal(gps_data["GPSLongitude"], gps_data.get("GPSLongitudeRef", "E"))
            data["gps"] = {"lat": lat, "lon": lon}
            location = entity("location", f"{lat:.6f},{lon:.6f}", f"{lat:.6f},{lon:.6f}", 0.8, source="exif", lat=lat, lon=lon)
            entities.append(location)
            edges.append(edge(file_entity.key, location.key, "located_at", 0.75))
            ev.append(evidence("geolocation", 4.0, f"https://www.google.com/maps?q={lat},{lon}", f"GPS dari EXIF: {lat}, {lon}", entity_key=location.key))
        except Exception as exc:
            warnings.append(f"GPS EXIF tidak bisa didekode: {exc}")
    if not data["exif"]:
        warnings.append("Tidak ada metadata EXIF yang bisa dipakai.")
    try:
        import pytesseract

        text = pytesseract.image_to_string(image)[:5000]
        if text.strip():
            data["ocr"] = text.strip()
            ocr_entities = _entities_from_text(text, "ocr")
            entities.extend(ocr_entities)
            for item in ocr_entities:
                edges.append(edge(file_entity.key, item.key, "contains_entity", 0.55))
            ev.append(evidence("ocr", 2.0, "", "Teks diekstrak dari gambar", entity_key=file_entity.key))
    except Exception:
        warnings.append("OCR tidak tersedia. Install pytesseract dan Tesseract untuk ekstraksi teks.")
    return ok(data, entities=entities, edges=edges, evidence=ev, warnings=warnings)


def analyze_document(path: str, filename: str) -> ToolResult:
    file_type = detect_file_type(path, filename)
    digest = sha256_file(path)
    file_entity = entity("document", filename, digest, 0.9, source="upload", filename=filename, file_type=file_type)
    entities = [file_entity]
    edges = []
    ev = [evidence("file", 1.0, "", f"File {filename}, SHA256 {digest[:16]}", entity_key=file_entity.key)]
    warnings = []
    data = {"filename": filename, "sha256": digest, "type": file_type, "metadata": {}, "text_preview": "", "pages": None}
    text = ""
    if file_type == "pdf":
        if PdfReader is None:
            return fail("dependency_missing", "pypdf tidak terpasang.")
        try:
            reader = PdfReader(path)
            data["pages"] = len(reader.pages)
            metadata = reader.metadata or {}
            for key, value in metadata.items():
                data["metadata"][str(key)] = str(value)[:300]
            for page in reader.pages[:10]:
                text += (page.extract_text() or "") + "\n"
        except Exception as exc:
            return fail("invalid_file", f"PDF tidak bisa dibaca: {exc}")
    elif file_type == "docx":
        if Document is None:
            return fail("dependency_missing", "python-docx tidak terpasang.")
        try:
            document = Document(path)
            props = document.core_properties
            data["metadata"] = {
                "author": props.author or "",
                "last_modified_by": props.last_modified_by or "",
                "title": props.title or "",
                "subject": props.subject or "",
                "created": str(props.created or ""),
                "modified": str(props.modified or ""),
            }
            text = "\n".join(paragraph.text for paragraph in document.paragraphs[:200])
        except Exception as exc:
            return fail("invalid_file", f"DOCX tidak bisa dibaca: {exc}")
    elif file_type == "xlsx":
        if load_workbook is None:
            return fail("dependency_missing", "openpyxl tidak terpasang.")
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            rows = []
            for sheet in workbook.worksheets[:3]:
                for row in sheet.iter_rows(values_only=True):
                    rows.append([str(cell) if cell is not None else "" for cell in row])
            text = "\n".join("\t".join(row) for row in rows[:500])
        except Exception as exc:
            return fail("invalid_file", f"XLSX tidak bisa dibaca: {exc}")
    elif file_type in {"csv", "text"}:
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as handle:
                text = handle.read(200000)
        except Exception as exc:
            return fail("invalid_file", f"File teks tidak bisa dibaca: {exc}")
    elif file_type == "json":
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as handle:
                payload = json.load(handle)
            text = json.dumps(payload, ensure_ascii=False)[:200000]
        except Exception as exc:
            return fail("invalid_file", f"JSON tidak bisa dibaca: {exc}")
    else:
        return fail("unsupported", "Format file belum didukung.")

    data["text_preview"] = re.sub(r"\s+", " ", text)[:800]
    author = data["metadata"].get("author") or data["metadata"].get("Author") or ""
    if author:
        person = entity("person", author, author.lower(), 0.6, source="document_metadata")
        entities.append(person)
        edges.append(edge(file_entity.key, person.key, "authored_by", 0.7))
        ev.append(evidence("document_metadata", 3.0, "", f"Author metadata: {author}", entity_key=person.key))
    extracted = _entities_from_text(text, "document_text")
    for item in extracted[:100]:
        entities.append(item)
        edges.append(edge(file_entity.key, item.key, "contains_entity", 0.6))
    if extracted:
        ev.append(evidence("document_text", 2.0, "", f"{len(extracted)} entity ditemukan di teks dokumen", entity_key=file_entity.key))
    if not extracted and not author:
        warnings.append("Tidak ada entity yang bisa diekstrak dari dokumen.")
    return ok(data, entities=entities, edges=edges, evidence=ev, warnings=warnings)


def analyze_structured(path: str, filename: str) -> ToolResult:
    file_type = detect_file_type(path, filename)
    if file_type == "csv":
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as handle:
            sample = handle.read(4096)
            handle.seek(0)
            dialect = csv.Sniffer().sniff(sample) if sample else csv.excel
            reader = csv.DictReader(handle, dialect=dialect)
            rows = [row for _, row in zip(range(200), reader)]
    elif file_type == "json":
        with open(path, "r", encoding="utf-8", errors="replace") as handle:
            payload = json.load(handle)
        if isinstance(payload, list):
            rows = [row if isinstance(row, dict) else {"value": row} for row in payload[:200]]
        elif isinstance(payload, dict):
            rows = [payload]
        else:
            rows = [{"value": payload}]
    else:
        return analyze_document(path, filename)
    result = analyze_document(path, filename)
    if result.status == "error":
        return result
    result.data["rows"] = rows[:50]
    result.data["row_count"] = len(rows)
    result.data["columns"] = sorted({key for row in rows for key in row.keys()})[:40]
    result.evidence.append(
        evidence(
            "structured_data",
            2.0,
            "",
            f"{len(rows)} baris dibaca dari {filename}",
            entity_key=f"document:{result.data.get('sha256', '')}",
        )
    )
    return result


def ingest_file(path: str, filename: str) -> ToolResult:
    file_type = detect_file_type(path, filename)
    if file_type == "image":
        return analyze_image(path, filename)
    if file_type in {"pdf", "docx", "xlsx", "text"}:
        return analyze_document(path, filename)
    if file_type in {"csv", "json"}:
        return analyze_structured(path, filename)
    return fail("unsupported", f"Format {file_type} belum didukung. Gunakan gambar, PDF, DOCX, XLSX, CSV, JSON, atau TXT.")
